import docx
from typing import Dict, List, Any
from app.core.errors import ParseException

def parse_docx(file_path: str) -> Dict[str, Any]:
    """
    Extracts structured text and layout metadata from a Word (.docx) document.
    """
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise ParseException(f"Failed to open DOCX document: {str(e)}")

    lines: List[str] = []
    has_tables = len(doc.tables) > 0

    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                lines.append(" | ".join(row_cells))

    combined_text = "\n".join(lines).strip()
    if not combined_text:
        raise ParseException("The DOCX document contains no readable text.")

    return {
        "raw_text": combined_text,
        "page_count": max(1, len(lines) // 40),
        "pages": [{"page_number": 1, "text": combined_text}],
        "parseability_flags": {
            "tables_detected": has_tables,
            "multi_columns_detected": False,
            "non_standard_fonts": False,
            "unusual_symbols": False,
        }
    }
